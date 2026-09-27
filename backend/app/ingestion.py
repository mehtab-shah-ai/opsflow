"""Bounded parsers. Structural uncertainty is surfaced instead of dropping records."""

import csv
import io
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse

MAX_ROWS = 20000
MAX_COLS = 100
MAX_CELLS = 500000
MAX_PAGES = 30


class IngestionError(ValueError):
    pass


def scalar(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def table_from_matrix(matrix, name, page=None):
    if not matrix:
        return None
    if len(matrix) > MAX_ROWS + 25:
        raise IngestionError(
            f"This table exceeds the {MAX_ROWS:,}-row limit. Split it into smaller files."
        )
    width = max((len(r) for r in matrix), default=0)
    if width > MAX_COLS or len(matrix) * width > MAX_CELLS:
        raise IngestionError(
            "This table exceeds the 100-column or 500,000-cell limit. Split it first."
        )
    # Skip title/blank lines only before a plausible multi-column header.
    candidates = [
        i
        for i, row in enumerate(matrix[:25])
        if sum(v not in (None, "") for v in row) >= min(2, width)
    ]
    if not candidates:
        return None
    header_idx = candidates[0]
    raw_headers = matrix[header_idx]
    headers, used, warnings = [], set(), []
    for i in range(width):
        original = raw_headers[i] if i < len(raw_headers) else None
        base = (
            re.sub(r"[^\w]+", "_", str(original or f"column_{i + 1}").strip().lower()).strip("_")
            or f"column_{i + 1}"
        )
        key, suffix = base, 2
        while key in used:
            key, suffix = f"{base}_{suffix}", suffix + 1
        if key != base:
            warnings.append(f"Duplicate header {base} retained as {key}.")
        used.add(key)
        headers.append(key)
    rows = []
    for source_row, row in enumerate(matrix[header_idx + 1 :], header_idx + 2):
        if not any(v not in (None, "") for v in row):
            continue
        if any(len(str(v)) > 10000 for v in row if v is not None):
            raise IngestionError(
                "A cell exceeds 10,000 characters. Shorten the cell before importing."
            )
        record = {col: scalar(row[i]) if i < len(row) else None for i, col in enumerate(headers)}
        record["_source_row"] = source_row
        if page:
            record["_source_page"] = page
        rows.append(record)
    if len(rows) > MAX_ROWS:
        raise IngestionError(f"This table exceeds the {MAX_ROWS:,}-row limit.")
    if not rows:
        return None
    if header_idx:
        warnings.append(
            f"Header detected at source row {header_idx + 1}; preceding title rows remain in the original file."
        )
    return {
        "name": name,
        "page": page,
        "rows": rows,
        "columns": headers,
        "warnings": warnings,
        "confidence": "High",
    }


def parse_file(filename: str, data: bytes) -> list[dict]:
    if not data:
        raise IngestionError(
            "This file is empty. Upload a file containing a header and records. No data was changed."
        )
    suffix = Path(filename).suffix.lower()
    tables = []
    try:
        if suffix == ".csv":
            if b"\x00" in data[:1000] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
                raise IngestionError(
                    "This does not look like a text CSV. Check the file extension."
                )
            if data.startswith((b"PK", b"%PDF", b"\xd0\xcf")):
                raise IngestionError("The file contents do not match its CSV extension.")
            text = None
            encodings = (
                ("utf-16",)
                if data.startswith((b"\xff\xfe", b"\xfe\xff"))
                else ("utf-8-sig", "cp1252")
            )
            for encoding in encodings:
                try:
                    text = data.decode(encoding)
                    break
                except UnicodeError:
                    continue
            if text is None:
                raise IngestionError("Unsupported text encoding. Save the file as UTF-8 CSV.")
            try:
                dialect = csv.Sniffer().sniff(text[:16000], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            matrix = []
            for row in csv.reader(io.StringIO(text), dialect, strict=True):
                matrix.append(row)
                if len(matrix) > MAX_ROWS + 25:
                    raise IngestionError("CSV exceeds the row limit. Split it into smaller files.")
            if matrix and any(len(r) > len(matrix[0]) for r in matrix[1:]):
                raise IngestionError(
                    "Some rows contain more fields than the header. Repair the delimiter or quoting; no fields were discarded."
                )
            tables = [table_from_matrix(matrix, "CSV data")]
        elif suffix == ".xlsx":
            if not data.startswith(b"PK"):
                raise IngestionError(
                    "This workbook is corrupt or password protected. Upload an unlocked XLSX copy."
                )
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 60 * 1024 * 1024:
                    raise IngestionError(
                        "Workbook expands beyond the 60 MB safety limit. Split it first."
                    )
                if any("vbaProject" in i.filename for i in archive.infolist()):
                    raise IngestionError(
                        "Macro-containing workbooks are not supported. Export values to XLSX or CSV."
                    )
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                if len(book.worksheets) > 20:
                    raise IngestionError("Workbook exceeds the 20-sheet limit.")
                total = 0
                for sheet in book.worksheets:
                    if sheet.max_row > MAX_ROWS + 25 or sheet.max_column > MAX_COLS:
                        raise IngestionError(
                            "Worksheet dimensions exceed safe limits. Remove unused formatted rows/columns or split the file."
                        )
                    total += sheet.max_row * sheet.max_column
                    if total > MAX_CELLS:
                        raise IngestionError("Workbook exceeds the 500,000-cell limit.")
                    matrix = [list(r) for r in sheet.iter_rows(values_only=True)]
                    tables.append(table_from_matrix(matrix, sheet.title))
            finally:
                book.close()
        elif suffix == ".xls":
            if not data.startswith(b"\xd0\xcf\x11\xe0"):
                raise IngestionError("Contents do not match a legacy XLS workbook.")
            import xlrd

            book = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                if book.nsheets > 20:
                    raise IngestionError("Workbook exceeds 20 sheets.")
                total = 0
                for sheet in book.sheets():
                    total += sheet.nrows * sheet.ncols
                    if sheet.nrows > MAX_ROWS + 25 or sheet.ncols > MAX_COLS or total > MAX_CELLS:
                        raise IngestionError("Legacy workbook exceeds safe dimensions.")
                    tables.append(
                        table_from_matrix(
                            [sheet.row_values(i) for i in range(sheet.nrows)], sheet.name
                        )
                    )
            finally:
                book.release_resources()
        elif suffix == ".pdf":
            if not data.startswith(b"%PDF"):
                raise IngestionError("Contents do not match a PDF document.")
            import pdfplumber

            with pdfplumber.open(io.BytesIO(data)) as pdf:
                if len(pdf.pages) > MAX_PAGES:
                    raise IngestionError(
                        f"PDF exceeds {MAX_PAGES} pages. Upload a smaller section."
                    )
                text_length = 0
                total_cells = 0
                for n, page in enumerate(pdf.pages, 1):
                    text_length += len(page.extract_text() or "")
                    for matrix in page.extract_tables():
                        total_cells += sum(len(r) for r in matrix)
                        if total_cells > MAX_CELLS:
                            raise IngestionError("PDF tables exceed the cell limit.")
                        tables.append(
                            table_from_matrix(matrix, f"Page {n} · Table {len(tables) + 1}", n)
                        )
                if text_length < 20:
                    raise IngestionError(
                        "This PDF appears scanned or empty. OCR/vision extraction is unavailable in this version. Upload a text PDF or spreadsheet. No data was changed."
                    )
                if not any(tables):
                    raise IngestionError(
                        "Text was found, but no reliable table could be detected. Export the table to CSV or Excel."
                    )
        else:
            raise IngestionError("Unsupported file type. Use CSV, XLSX, XLS or a native PDF.")
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError(
            "The file could not be read safely. It may be corrupt, encrypted, or malformed. Upload an unlocked, valid copy. No data was changed."
        ) from exc
    tables = [t for t in tables if t]
    if not tables:
        raise IngestionError("No data rows were found. Include a header and at least one record.")
    return tables


def sheet_csv_url(url: str) -> str:
    parsed = urlparse(url)
    match = re.fullmatch(r"/spreadsheets/d/([A-Za-z0-9_-]+)/?(?:edit|view|preview)?/?", parsed.path)
    if parsed.scheme != "https" or parsed.netloc != "docs.google.com" or not match:
        raise IngestionError("Use a public https://docs.google.com/spreadsheets/d/... URL.")
    gid = parse_qs(parsed.fragment).get("gid", parse_qs(parsed.query).get("gid", ["0"]))[0]
    if not gid.isdigit():
        raise IngestionError("Worksheet gid must be numeric.")
    return f"https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=csv&gid={gid}"
