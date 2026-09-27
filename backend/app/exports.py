import csv
import io
from datetime import datetime, timezone
from xml.sax.saxutils import escape

from .domain import columns_of


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + value
    return value


def export_table(rows, fmt, columns=None):
    columns = columns or columns_of(rows) or ["No_records"]
    if fmt == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow([safe_cell(c) for c in columns])
        writer.writerows([[safe_cell(r.get(c)) for c in columns] for r in rows])
        return stream.getvalue().encode("utf-8-sig")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "OpsFlow export"
    sheet.append([safe_cell(c) for c in columns])
    for row in rows:
        sheet.append([safe_cell(row.get(c)) for c in columns])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="173F35")
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        sheet.column_dimensions[cells[0].column_letter].width = min(
            42, max(16, len(str(cells[0].value)) + 4)
        )
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def report_pdf(dataset, audit, report="daily", period="all"):
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (
        KeepTogether,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            "Brand",
            fontName="Helvetica-Bold",
            fontSize=27,
            leading=33,
            textColor=colors.HexColor("#173F35"),
            spaceAfter=16,
        )
    )
    styles.add(
        ParagraphStyle("Muted", fontSize=9, leading=14, textColor=colors.HexColor("#64756C"))
    )
    styles.add(ParagraphStyle("Cell", fontSize=8, leading=11))
    styles.add(ParagraphStyle("Right", fontSize=8, alignment=TA_RIGHT))
    title = {
        "daily": "Operations MIS Report",
        "quality": "Data Quality & Health Report",
        "exceptions": "Operational Exceptions & Issues Report",
    }[report]
    a = dataset["analysis"]
    story = [
        Paragraph("OpsFlow AI", styles["Brand"]),
        Paragraph(title, styles["Heading1"]),
        Paragraph(
            escape(f"{dataset['filename']} | Version {dataset['version']} | Period: {period}"),
            styles["Muted"],
        ),
        Paragraph(
            "Synthetic demo dataset"
            if dataset.get("demo")
            else "User-supplied operational dataset",
            styles["Muted"],
        ),
        Spacer(1, 18),
    ]

    def table(data, widths=None):
        rendered = [[Paragraph(escape(str(v)), styles["Cell"]) for v in row] for row in data]
        t = Table(rendered, colWidths=widths, repeatRows=1, hAlign="LEFT")
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E4F2E9")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DEE5DF")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            )
        )
        return t

    story += [
        table(
            [
                ["Total Records", "Initial Quality", "Current Quality", "Items Flagged"],
                [
                    a["rows"],
                    f"{dataset['quality_before']}%",
                    f"{a['quality_score']}%",
                    len(a["issues"]),
                ],
            ],
            [120] * 4,
        ),
        Spacer(1, 18),
        Paragraph("Executive Summary", styles["Heading2"]),
        Paragraph(escape(a["summary"]), styles["BodyText"]),
        Spacer(1, 14),
    ]
    if a["operations"]["available"]:
        op = a["operations"]
        story += [
            Paragraph("Key Operational Indicators (KPIs)", styles["Heading2"]),
            table(
                [
                    ["Staff Attendance", "Shift Coverage", "Task Completion", "SLA On-Time Rate"],
                    *[
                        [
                            f"{op[k]}%"
                            for k in (
                                "attendance_rate",
                                "shift_coverage",
                                "task_completion",
                                "sla_compliance",
                            )
                        ]
                    ],
                ],
                [120] * 4,
            ),
            Spacer(1, 16),
        ]
        bars = Drawing(480, 35 + len(op["sites"][:8]) * 28)
        scale = max([s["required"] for s in op["sites"]] or [1]) or 1
        for i, s in enumerate(op["sites"][:8]):
            y = bars.height - 28 - i * 28
            bars.add(String(0, y + 4, s["site"][:24], fontSize=9))
            bars.add(
                Rect(
                    120,
                    y,
                    280 * s["available"] / scale,
                    12,
                    fillColor=colors.HexColor("#217A65"),
                    strokeColor=None,
                )
            )
            bars.add(String(410, y + 3, f"{s['available']:.0f}/{s['required']:.0f}", fontSize=8))
        story += [
            KeepTogether(
                [Paragraph("Site Staffing Coverage (Present vs Required Staff)", styles["Heading3"]), bars]
            ),
            Spacer(1, 12),
        ]
        top = op["sites"][0] if op["sites"] else None
        if top:
            story.append(
                Paragraph(
                    escape(
                        f"Recommended Priority Action: Review staffing coverage at {top['site']}. There are {top['gap']:g} uncovered shifts that require replacement staff to prevent operational delays."
                    ),
                    styles["BodyText"],
                )
            )
    story += [Spacer(1, 16), Paragraph("Priority Issues & Flagged Items", styles["Heading2"])]
    ordered = sorted(
        a["issues"], key=lambda i: {"critical": 0, "warning": 1, "info": 2}[i["severity"]]
    )
    severity_label = {"critical": "Urgent", "warning": "Review", "info": "Minor"}
    story += [
        table(
            [["Row", "Status", "Field Name", "What Was Found & Recommended Action"]]
            + [[i["row"], severity_label.get(i["severity"], i["severity"]), i["column"], i["reason"]] for i in ordered[:30]],
            [40, 55, 95, 290],
        ),
        Paragraph(
            f"Showing top {min(30, len(ordered))} of {len(ordered)} flagged items. For the complete list, download the Excel exceptions file.",
            styles["Muted"],
        ),
        Spacer(1, 14),
        Paragraph("Data Safety & Traceability", styles["Heading2"]),
        Paragraph(
            escape(
                f"{len(audit)} approved changes applied so far. {a['safe_changes']} safe formatting corrections are ready for review. No original records were deleted or lost. Your original uploaded data is completely preserved."
            ),
            styles["BodyText"],
        ),
        Spacer(1, 12),
        Paragraph(
            "Quality Score Guide: Starts at 100% (clean data). Missing values and duplicate shifts lower the score the most, while simple formatting cleanups have minor impact. This score measures data accuracy and completeness.",
            styles["Muted"],
        ),
        Spacer(1, 12),
        Paragraph(
            escape(
                f"Dataset: {dataset['filename']} | Run ID: {dataset['job_id'][:12]} | Generated: {datetime.now(timezone.utc).strftime('%d %b %Y, %I:%M %p UTC')}"
            ),
            styles["Muted"],
        ),
    ]

    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64756C"))
        canvas.drawString(48, 26, "OPSFLOW AI  /  INTELLIGENT BUSINESS OPERATIONS")
        canvas.drawRightString(545, 26, str(doc.page))

    SimpleDocTemplate(buf, rightMargin=48, leftMargin=48, topMargin=40, bottomMargin=44).build(
        story, onFirstPage=footer, onLaterPages=footer
    )
    return buf.getvalue()
